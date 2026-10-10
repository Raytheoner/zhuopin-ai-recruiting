"""文书引擎端到端端点（U2 tasks 2.1/2.4/2.5/2.6）。"""
from __future__ import annotations

import io
import json
import sqlite3

import docx
from fastapi.testclient import TestClient

from app.agents.jd_agent import AI_LABEL_PREFIX
from app.audit.hook import RecorderAuditHook
from app.audit.recorder import AuditRecorder
from app.audit.sinks import JsonlChainSink, SqliteSink
from app.llm.gateway import LLMGateway
from app.storage.auth_session import create_session
from app.storage.db import get_connection, init_schema
from app.storage.hr_account import upsert_account
from app.web.server import create_app
from tests.test_web_api import ScriptedOpenAIClient


def _app(tmp_path, responses):
    db_path = str(tmp_path / "web.db")
    conn = get_connection(db_path)
    init_schema(conn)
    recorder = AuditRecorder(SqliteSink(conn), JsonlChainSink(tmp_path / "decisions.jsonl"))
    hook = RecorderAuditHook(recorder, conn)
    scripted = ScriptedOpenAIClient(responses)

    def gateway_factory():
        return LLMGateway(
            api_key="k", base_url="https://example.com", model="deepseek-chat",
            supports_json_schema=False, client=scripted, audit_hook=hook,
        )

    app = create_app(db_path=db_path, gateway_factory=gateway_factory, root_path="")
    return TestClient(app), scripted


def _login(tmp_path, client, username="tester"):
    conn = sqlite3.connect(str(tmp_path / "web.db"))
    account_id = upsert_account(conn, username=username, password="testpass123")
    token = create_session(conn, hr_account_id=account_id)
    conn.close()
    client.cookies.set("hr_session", token)


def _seed_app(tmp_path):
    conn = sqlite3.connect(str(tmp_path / "web.db"))
    conn.execute("INSERT INTO job (id, title, status) VALUES ('j1','嵌入式工程师','approved')")
    conn.execute("INSERT INTO candidate (id, name) VALUES ('c1','张三')")
    conn.execute(
        "INSERT INTO resume (id, job_id, sample_class, file_name, content_sha256, "
        "status, uploaded_by) VALUES ('r1','j1','synthetic','a.pdf','sha','parsed','hr')"
    )
    conn.execute(
        "INSERT INTO application (id, candidate_id, job_id, resume_id, current_stage_id) "
        "VALUES ('app1','c1','j1','r1','offer')"
    )
    conn.execute(
        "INSERT INTO offer (id, application_id, job_id, department, start_date, "
        "report_to, status, created_by) VALUES ('o1','app1','j1','研发部','2026-10-20',"
        "'李四','approved','hr')"
    )
    conn.commit()
    conn.close()


def test_template_put_rejects_salary_placeholder(tmp_path):
    client, _ = _app(tmp_path, [])
    _login(tmp_path, client)
    resp = client.put("/api/letter-templates/offer", json={"body": "{salary} 您好"})
    assert resp.status_code == 422
    assert "salary" in resp.json()["detail"]


def test_generate_twice_creates_new_version(tmp_path):
    client, _ = _app(tmp_path, [
        json.dumps({"body": "拟录用您担任嵌入式工程师"}),
        json.dumps({"body": "拟录用您担任嵌入式工程师"}),
    ])
    _login(tmp_path, client)
    _seed_app(tmp_path)

    resp = client.post("/api/applications/app1/letters", json={"kind": "offer"})
    assert resp.status_code == 201
    letter = resp.json()
    assert letter["version"] == 1
    assert letter["ai_generated"] is True
    assert AI_LABEL_PREFIX in letter["body"]
    # 同一投递重复生成 MUST 产生新版本（spec Scenario「生成 Offer 草稿」的
    # 「同一投递重复生成」）。两次生成的请求标识 nonce 不同 ⇒ input_hash 不同 ⇒
    # 两个不同的 run_id ⇒ 两行 candidate_letter。
    resp2 = client.post("/api/applications/app1/letters", json={"kind": "offer"})
    assert resp2.status_code == 201
    assert resp2.json()["version"] == 2


def test_generate_rejection_requires_rejection_record(tmp_path):
    client, _ = _app(tmp_path, [])
    _login(tmp_path, client)
    _seed_app(tmp_path)
    resp = client.post("/api/applications/app1/letters", json={"kind": "rejection"})
    assert resp.status_code == 409


def test_edit_preserves_ai_label_and_mark_human_strips_it(tmp_path):
    client, _ = _app(tmp_path, [json.dumps({"body": "拟录用您担任嵌入式工程师"})])
    _login(tmp_path, client)
    _seed_app(tmp_path)
    letter = client.post("/api/applications/app1/letters", json={"kind": "offer"}).json()

    edited = client.patch(f"/api/letters/{letter['id']}", json={"body": "改过的正文"}).json()
    assert edited["ai_generated"] is True
    assert AI_LABEL_PREFIX in edited["body"]

    marked = client.post(f"/api/letters/{letter['id']}/mark-human").json()
    assert marked["ai_generated"] is False
    assert AI_LABEL_PREFIX not in marked["body"]
    assert marked["authorship_marked_by"] == "tester"
    assert marked["authorship_from_version"] == 1


def test_view_writes_access_log_then_returns_body(tmp_path):
    client, _ = _app(tmp_path, [json.dumps({"body": "拟录用您担任嵌入式工程师"})])
    _login(tmp_path, client)
    _seed_app(tmp_path)
    letter = client.post("/api/applications/app1/letters", json={"kind": "offer"}).json()
    resp = client.get(f"/api/letters/{letter['id']}")
    assert resp.status_code == 200
    assert resp.json()["body"] == letter["body"]
    conn = sqlite3.connect(str(tmp_path / "web.db"))
    row = conn.execute(
        "SELECT access_type, accessor FROM letter_access_log WHERE letter_id = ?",
        (letter["id"],),
    ).fetchone()
    assert row == ("view", "tester")


def test_view_log_failure_blocks_body(tmp_path, monkeypatch):
    client, _ = _app(tmp_path, [json.dumps({"body": "拟录用您担任嵌入式工程师"})])
    _login(tmp_path, client)
    _seed_app(tmp_path)
    letter = client.post("/api/applications/app1/letters", json={"kind": "offer"}).json()

    def boom(*args, **kwargs):
        raise RuntimeError("留痕表不可写")

    import app.web.server as server_mod
    monkeypatch.setattr(server_mod, "record_letter_access", boom)
    # ⚠️ 默认 TestClient(raise_server_exceptions=True) 会把 ServerErrorMiddleware
    # 已经发出的 500 响应吞掉、转成 pytest 里的一次异常重新抛出。本用例断言的是
    # 「路由返回了 500」，不是「异常发生过」——同 tests/test_resume_access_log.py::
    # test_access_log_write_failure_blocks_the_read 的既有做法：用同一个 app 另建
    # 一个 raise_server_exceptions=False 的客户端，带上原客户端的登录态 cookie。
    no_raise_client = TestClient(client.app, raise_server_exceptions=False)
    no_raise_client.cookies.update(client.cookies)
    resp = no_raise_client.get(f"/api/letters/{letter['id']}")
    assert resp.status_code == 500


def test_export_writes_log_sets_sent_status_and_renders(tmp_path):
    client, _ = _app(tmp_path, [json.dumps({"body": "拟录用您担任嵌入式工程师"})])
    _login(tmp_path, client)
    _seed_app(tmp_path)
    letter = client.post("/api/applications/app1/letters", json={"kind": "offer"}).json()

    resp = client.get(f"/api/letters/{letter['id']}/export.docx")
    assert resp.status_code == 200
    document = docx.Document(io.BytesIO(resp.content))
    header = "\n".join(p.text for p in document.sections[0].header.paragraphs)
    assert AI_LABEL_PREFIX in header
    texts = [p.text for p in document.paragraphs]
    assert "薪资待遇（由 HR 手工填写）：" in texts

    conn = sqlite3.connect(str(tmp_path / "web.db"))
    assert conn.execute(
        "SELECT sent_status FROM candidate_letter WHERE id = ?", (letter["id"],)
    ).fetchone()[0] == "exported"
    assert conn.execute(
        "SELECT COUNT(*) FROM letter_access_log WHERE letter_id = ? AND access_type = 'export'",
        (letter["id"],),
    ).fetchone()[0] == 1
