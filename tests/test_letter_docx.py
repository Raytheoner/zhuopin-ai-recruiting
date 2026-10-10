"""letter_docx 渲染（U2 tasks 2.5）。"""
from __future__ import annotations

import docx
from docx.oxml.ns import qn

from app.agents.jd_agent import AI_LABEL_PREFIX, AI_LABEL_TEMPLATE
from app.letter_docx import (
    DEFAULT_BODY_FONT,
    SALARY_PARAGRAPH_LABEL,
    render_letter_to_docx,
)

BODY = (
    "张三：\n\n经评估，拟录用您担任嵌入式工程师。\n\n"
    + AI_LABEL_TEMPLATE.format(generated_at="2026-10-10 00:00:00")
)


def test_offer_docx_header_contains_ai_label_when_ai_generated(tmp_path):
    out = render_letter_to_docx(
        body=BODY, kind="offer", ai_generated=True, out_path=tmp_path / "a.docx"
    )
    document = docx.Document(str(out))
    header = "\n".join(p.text for p in document.sections[0].header.paragraphs)
    assert AI_LABEL_PREFIX in header


def test_human_written_docx_header_has_no_ai_label(tmp_path):
    out = render_letter_to_docx(
        body="张三：您好", kind="rejection", ai_generated=False,
        out_path=tmp_path / "a.docx",
    )
    document = docx.Document(str(out))
    header = "\n".join(p.text for p in document.sections[0].header.paragraphs)
    assert AI_LABEL_PREFIX not in header


def test_offer_salary_paragraph_is_blank(tmp_path):
    out = render_letter_to_docx(
        body=BODY, kind="offer", ai_generated=True, out_path=tmp_path / "a.docx"
    )
    document = docx.Document(str(out))
    texts = [p.text for p in document.paragraphs]
    assert SALARY_PARAGRAPH_LABEL in texts
    assert texts[texts.index(SALARY_PARAGRAPH_LABEL) + 1] == ""


def test_rejection_has_no_salary_paragraph(tmp_path):
    out = render_letter_to_docx(
        body=BODY, kind="rejection", ai_generated=False, out_path=tmp_path / "a.docx"
    )
    document = docx.Document(str(out))
    assert all(SALARY_PARAGRAPH_LABEL not in p.text for p in document.paragraphs)


def test_body_text_preserved_and_runs_pin_east_asia_font(tmp_path):
    out = render_letter_to_docx(
        body=BODY, kind="offer", ai_generated=False, out_path=tmp_path / "a.docx"
    )
    document = docx.Document(str(out))
    texts = [p.text for p in document.paragraphs]
    assert any("嵌入式工程师" in t for t in texts)
    fonts = set()
    for p in document.paragraphs:
        for run in p.runs:
            rpr = run._element.rPr
            if rpr is not None and rpr.rFonts is not None:
                fonts.add(rpr.rFonts.get(qn("w:eastAsia")))
    assert fonts <= {DEFAULT_BODY_FONT, "Consolas"}
