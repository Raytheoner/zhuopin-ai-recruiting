"""scripts/letter_md_to_docx.py：字体统一（东亚字体必须显式钉死）与结构渲染。"""
from __future__ import annotations

import docx
from docx.oxml.ns import qn

from scripts.letter_md_to_docx import DEFAULT_BODY_FONT, render_letter_docx

SAMPLE = """---
title: 示例信
---
# 示例#1 · 标题

你好：

## 一、小节

- **要点**：带 `code` 的一行
1. 第一项

普通段落。
"""


def _east_asia(style) -> str | None:
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    return rfonts.get(qn("w:eastAsia")) if rfonts is not None else None


def test_every_style_pins_the_same_east_asia_font(tmp_path) -> None:
    src = tmp_path / "a.md"
    src.write_text(SAMPLE, encoding="utf-8")
    out = render_letter_docx(src, tmp_path / "a.docx")
    document = docx.Document(str(out))
    for name in ("Normal", "Heading 1", "Heading 2", "Heading 3", "List Bullet", "List Number"):
        assert _east_asia(document.styles[name]) == DEFAULT_BODY_FONT, name


def test_runs_do_not_carry_conflicting_fonts(tmp_path) -> None:
    src = tmp_path / "a.md"
    src.write_text(SAMPLE, encoding="utf-8")
    out = render_letter_docx(src, tmp_path / "a.docx")
    document = docx.Document(str(out))
    fonts = set()
    for paragraph in document.paragraphs:
        for run in paragraph.runs:
            rpr = run._element.rPr
            if rpr is not None and rpr.rFonts is not None:
                fonts.add(rpr.rFonts.get(qn("w:eastAsia")))
    assert fonts <= {DEFAULT_BODY_FONT, "Consolas"}, fonts


def test_structure_rendered(tmp_path) -> None:
    src = tmp_path / "a.md"
    src.write_text(SAMPLE, encoding="utf-8")
    out = render_letter_docx(src, tmp_path / "a.docx")
    document = docx.Document(str(out))
    styles = [p.style.name for p in document.paragraphs]
    assert "Heading 1" in styles and "Heading 2" in styles
    assert "List Bullet" in styles and "List Number" in styles
    # frontmatter 不进入正文
    assert all("title:" not in p.text for p in document.paragraphs)
