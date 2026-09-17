from pathlib import Path

import pytest
from PIL import Image
from reportlab.pdfgen import canvas

from app.parsing.extract_text import (
    MIN_EFFECTIVE_CHARS,
    OcrUnavailable,
    UnsupportedFileType,
    effective_char_count,
    extract_text,
)

LONG_ASCII = "Resume of Test Person. Skills: C, AUTOSAR CP, CAN, LIN, UDS, Bootloader, MCAL, ISO 26262 functional safety."


def _text_pdf(path: Path, text: str) -> Path:
    c = canvas.Canvas(str(path))
    c.setFont("Helvetica", 12)
    c.drawString(72, 720, text)
    c.save()
    return path


def _scan_pdf(path: Path) -> Path:
    Image.new("RGB", (600, 200), "white").save(str(path), "PDF")
    return path


def _docx(path: Path, lines: list[str]) -> Path:
    import docx

    document = docx.Document()
    for line in lines:
        document.add_paragraph(line)
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "技能"
    table.rows[0].cells[1].text = "C、AUTOSAR"
    document.save(str(path))
    return path


def test_effective_char_count_ignores_whitespace():
    assert effective_char_count(" a b\n c ") == 3
    assert MIN_EFFECTIVE_CHARS == 50


def test_text_pdf_is_extracted_directly_without_ocr(tmp_path):
    def boom(_):
        raise AssertionError("文本型 PDF 不该走 OCR")

    out = extract_text(_text_pdf(tmp_path / "a.pdf", LONG_ASCII), ocr=boom)
    assert out.kind == "pdf_text"
    assert out.readable is True
    assert "AUTOSAR" in out.text


def test_scan_pdf_routes_to_injected_ocr(tmp_path):
    out = extract_text(_scan_pdf(tmp_path / "s.pdf"), ocr=lambda _: LONG_ASCII)
    assert out.kind == "pdf_scan"
    assert out.readable is True
    assert out.text == LONG_ASCII


def test_scan_pdf_with_empty_ocr_is_unreadable(tmp_path):
    out = extract_text(_scan_pdf(tmp_path / "s.pdf"), ocr=lambda _: "  \n ")
    assert out.kind == "pdf_scan"
    assert out.readable is False
    assert out.effective_chars == 0


def test_ocr_unavailable_propagates(tmp_path):
    def missing(_):
        raise OcrUnavailable("paddleocr 未安装")

    with pytest.raises(OcrUnavailable):
        extract_text(_scan_pdf(tmp_path / "s.pdf"), ocr=missing)


def test_docx_paragraphs_and_tables_are_joined(tmp_path):
    lines = ["张明远", "期望工作城市：无锡"] + ["填充行" * 10] * 3
    out = extract_text(_docx(tmp_path / "r.docx", lines))
    assert out.kind == "docx"
    assert out.readable is True
    assert "期望工作城市：无锡" in out.text
    assert "技能\tC、AUTOSAR" in out.text


def test_unsupported_suffix_rejected(tmp_path):
    p = tmp_path / "r.txt"
    p.write_text("x", encoding="utf-8")
    with pytest.raises(UnsupportedFileType):
        extract_text(p)


def test_default_ocr_raises_unavailable_when_paddle_missing(tmp_path, monkeypatch):
    import builtins

    real_import = builtins.__import__

    def fake_import(name, *args, **kwargs):
        if name in {"paddleocr", "pymupdf", "fitz"}:
            raise ImportError(f"No module named {name!r}")
        return real_import(name, *args, **kwargs)

    monkeypatch.setattr(builtins, "__import__", fake_import)
    with pytest.raises(OcrUnavailable):
        extract_text(_scan_pdf(tmp_path / "s.pdf"))
