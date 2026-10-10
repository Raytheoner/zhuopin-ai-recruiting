"""Offer/拒信文书 docx 渲染（offer-generation U2 tasks 2.5）。

字体统一做法与 scripts/letter_md_to_docx.py 一致（微软雅黑 + Consolas，东亚字体
必须显式钉死，否则中文各自回退导致参差不齐）——但本模块面向 candidate_letter.body
字符串渲染，并额外承担两件脚本不承担的事：
① 未标记人工撰写（ai_generated=1）时页眉写 AI 标识文字；
② Offer 类在正文后追加「薪资待遇（由 HR 手工填写）：」+ 一个空段落（薪资位置留空）。
⛔ 本模块不 import app.storage / scripts（app→scripts 是层次倒置）。
"""
from __future__ import annotations

import re
from pathlib import Path

import docx
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

from app.agents.jd_agent import (
    AI_LABEL_TEMPLATE,
    UNKNOWN_GENERATED_AT,
    extract_label_generated_at,
)

DEFAULT_BODY_FONT = "微软雅黑"
DEFAULT_MONO_FONT = "Consolas"
SALARY_PARAGRAPH_LABEL = "薪资待遇（由 HR 手工填写）："


def _pin_font(
    style,
    *,
    ascii_font: str,
    east_asia_font: str,
    size_pt: float,
    bold: bool | None = None,
    color: tuple[int, int, int] | None = None,
    space_after_pt: float | None = None,
) -> None:
    style.font.name = ascii_font
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = rpr.makeelement(qn("w:rFonts"), {})
        rpr.append(rfonts)
    rfonts.set(qn("w:ascii"), ascii_font)
    rfonts.set(qn("w:hAnsi"), ascii_font)
    rfonts.set(qn("w:eastAsia"), east_asia_font)
    style.font.size = Pt(size_pt)
    if bold is not None:
        style.font.bold = bold
    if color is not None:
        style.font.color.rgb = RGBColor(*color)
    if space_after_pt is not None:
        style.paragraph_format.space_after = Pt(space_after_pt)


def _add_runs(paragraph, text: str) -> None:
    for part in re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
            run.font.name = DEFAULT_BODY_FONT
            run._element.rPr.rFonts.set(qn("w:eastAsia"), DEFAULT_BODY_FONT)
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = DEFAULT_MONO_FONT
            run._element.rPr.rFonts.set(qn("w:eastAsia"), DEFAULT_MONO_FONT)
        else:
            run = paragraph.add_run(part)
            run.font.name = DEFAULT_BODY_FONT
            run._element.rPr.rFonts.set(qn("w:eastAsia"), DEFAULT_BODY_FONT)


def _pin_default_styles(document: "docx.Document") -> None:
    _pin_font(
        document.styles["Normal"], ascii_font=DEFAULT_BODY_FONT,
        east_asia_font=DEFAULT_BODY_FONT, size_pt=11, space_after_pt=6,
    )
    for name, size in (("Heading 1", 16), ("Heading 2", 13.5), ("Heading 3", 12)):
        _pin_font(
            document.styles[name], ascii_font=DEFAULT_BODY_FONT,
            east_asia_font=DEFAULT_BODY_FONT, size_pt=size, bold=True,
            color=(0x1F, 0x1F, 0x1F), space_after_pt=6,
        )
    for list_style in ("List Bullet", "List Number"):
        _pin_font(
            document.styles[list_style], ascii_font=DEFAULT_BODY_FONT,
            east_asia_font=DEFAULT_BODY_FONT, size_pt=11, space_after_pt=4,
        )


def _set_ai_header(document: "docx.Document", body: str) -> None:
    generated_at = extract_label_generated_at(body) or UNKNOWN_GENERATED_AT
    paragraph = document.sections[0].header.paragraphs[0]
    paragraph.text = AI_LABEL_TEMPLATE.format(generated_at=generated_at)
    for run in paragraph.runs:
        run.font.name = DEFAULT_BODY_FONT
        run._element.rPr.rFonts.set(qn("w:eastAsia"), DEFAULT_BODY_FONT)


def _append_body_lines(document: "docx.Document", body: str) -> None:
    for raw in body.splitlines():
        line = raw.rstrip()
        if not line.strip():
            continue
        if line.startswith("### "):
            document.add_heading(line[4:].strip(), level=3)
        elif line.startswith("## "):
            document.add_heading(line[3:].strip(), level=2)
        elif line.startswith("# "):
            document.add_heading(line[2:].strip(), level=1)
        elif re.match(r"^\s*[-*] ", line):
            _add_runs(
                document.add_paragraph(style="List Bullet"),
                re.sub(r"^\s*[-*] ", "", line),
            )
        elif re.match(r"^\s*\d+\. ", line):
            _add_runs(
                document.add_paragraph(style="List Number"),
                re.sub(r"^\s*\d+\. ", "", line),
            )
        else:
            _add_runs(document.add_paragraph(), line)


def render_letter_to_docx(
    *,
    body: str,
    kind: str,
    ai_generated: bool,
    out_path: str | Path,
) -> Path:
    """把一版文书草稿渲染成 docx。导出内容与草稿一致（body 原样渲染，含 AI 标识行），
    页眉按 ai_generated 决定是否写 AI 标识；Offer 末尾追加留空的薪资段落。"""
    document = docx.Document()
    _pin_default_styles(document)
    if ai_generated:
        _set_ai_header(document, body)
    _append_body_lines(document, body)
    if kind == "offer":
        _add_runs(document.add_paragraph(), SALARY_PARAGRAPH_LABEL)
        document.add_paragraph()  # 薪资位置：空段落，由 HR 手填
    out = Path(out_path)
    document.save(str(out))
    return out
