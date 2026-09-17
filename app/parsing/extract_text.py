"""
简历文件 → 文本（spec「扫描件与不可读文件」，design D14）。

文本型 PDF 直抽（pypdf）、Word 走 python-docx、扫描件走 PaddleOCR（懒加载；缺包抛
OcrUnavailable，由调用方决定进"不可读"队列还是报错）。OCR 可注入，单测不碰 Paddle。
"""
from __future__ import annotations

from collections.abc import Callable
from dataclasses import dataclass
from pathlib import Path
from typing import Literal

MIN_EFFECTIVE_CHARS = 50
SUPPORTED_SUFFIXES = (".pdf", ".docx")


class UnsupportedFileType(ValueError):
    """只收 pdf/docx（design D1 类型白名单）。"""


class OcrUnavailable(RuntimeError):
    """PaddleOCR / PyMuPDF 未安装或初始化失败。"""


@dataclass(frozen=True)
class ExtractedText:
    text: str
    kind: Literal["pdf_text", "pdf_scan", "docx"]
    effective_chars: int
    readable: bool


def effective_char_count(text: str) -> int:
    return sum(1 for ch in text if not ch.isspace())


def extract_docx(path: Path) -> str:
    import docx

    document = docx.Document(str(path))
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.append("\t".join(cell.text for cell in row.cells))
    return "\n".join(parts)


def extract_pdf_text(path: Path) -> str:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    return "\n".join((page.extract_text() or "") for page in reader.pages)


def _paddle_lines(result: object) -> list[str]:
    """兼容 PaddleOCR 2.x（[[box, (text, score)], ...]）与 3.x（含 rec_texts 的结果对象）两种返回形态。"""
    lines: list[str] = []
    for page in result or []:
        rec_texts = None
        if isinstance(page, dict):
            rec_texts = page.get("rec_texts")
        elif hasattr(page, "get"):
            try:
                rec_texts = page.get("rec_texts")
            except Exception:
                rec_texts = None
        if rec_texts is not None:
            lines.extend(str(t) for t in rec_texts)
            continue
        for item in page or []:
            try:
                lines.append(str(item[1][0]))
            except (TypeError, IndexError, KeyError):
                continue
    return lines


def ocr_pdf(path: Path, *, lang: str = "ch") -> str:
    """每次调用都新建一个 PaddleOCR 引擎（初始化耗时不小）。U2 的常驻消费方（tasks 3.4）应缓存该引擎实例，
    ⛔ 不要沿用这里"每次新建"的写法——本函数只服务离线一次性核对场景（M9）。"""
    try:
        import numpy as np
        from paddleocr import PaddleOCR

        try:
            import pymupdf as fitz  # 1.24+ 的正式模块名；`import fitz` 已标 deprecated
        except ImportError:
            import fitz
    except ImportError as exc:
        raise OcrUnavailable(f"PaddleOCR/PyMuPDF 未安装: {exc}") from exc
    try:
        engine = PaddleOCR(lang=lang)
    except Exception as exc:  # 模型下载失败、DLL 缺失等，都属"OCR 不可用"
        raise OcrUnavailable(f"PaddleOCR 初始化失败: {exc!r}") from exc
    lines: list[str] = []
    with fitz.open(str(path)) as doc:
        for page in doc:
            pix = page.get_pixmap(dpi=200)
            image = np.frombuffer(pix.samples, dtype=np.uint8).reshape(pix.height, pix.width, pix.n)
            if pix.n == 4:
                image = image[:, :, :3]
            result = engine.ocr(image)
            lines.extend(_paddle_lines(result))
    return "\n".join(lines)


def extract_text(
    path: Path,
    *,
    ocr: Callable[[Path], str] | None = None,
    min_effective_chars: int = MIN_EFFECTIVE_CHARS,
) -> ExtractedText:
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix == ".docx":
        text = extract_docx(path)
        kind: Literal["pdf_text", "pdf_scan", "docx"] = "docx"
    elif suffix == ".pdf":
        text = extract_pdf_text(path)
        kind = "pdf_text"
        if effective_char_count(text) < min_effective_chars:
            kind = "pdf_scan"
            text = (ocr or ocr_pdf)(path)
    else:
        raise UnsupportedFileType(f"只支持 {SUPPORTED_SUFFIXES}，收到 {path.name}")
    n = effective_char_count(text)
    return ExtractedText(text=text, kind=kind, effective_chars=n, readable=n >= min_effective_chars)
